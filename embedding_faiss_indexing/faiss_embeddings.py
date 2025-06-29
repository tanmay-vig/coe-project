from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import OllamaEmbeddings
from docx import Document
import json

# load questions
with open("questions.json") as f:
    questions = json.load(f)

#  word document
doc = Document()
for i, q in enumerate(questions):
    doc.add_paragraph(f"Q{i+1}: {q['question']}")
    doc.add_paragraph(f"Topic: {q['topic']} | Type: {q['type']} | "
                     f"Difficulty: {q['difficulty']} | Cognitive: {q['cognitive_level']}")
    if( q["type"] == "mcq"):
        doc.add_paragraph("Options:")
        for j, option in enumerate(q["options"], start=1):
            doc.add_paragraph(f"  {j}. {option}")
    doc.add_paragraph("\n")
doc.save("question_bank.docx")

# FAISS index
texts = [q["question"] for q in questions]
metadatas = [
    {
        "topic": q["topic"],
        "type": q["type"],
        "difficulty": q["difficulty"],
        "cognitive_level": q["cognitive_level"],
        **({"options": q["options"]} if q["type"] == "mcq" else {}) #dict unpacking
    }
    for q in questions
]


embeddings = OllamaEmbeddings(model="llama3.2")
vectorstore = FAISS.from_texts(texts, embeddings, metadatas=metadatas)
vectorstore.save_local("faiss_index")